
# -*- coding: utf-8 -*-
"""
BSS完整系统 - 自行车共享系统(Bike Sharing System)需求预测与车队优化

该系统实现了以下功能:
1. 需求预测: 基于历史数据预测各自行车站点的净需求
2. 车队优化: 使用传统算法和神经网络两种方式优化配车方案
3. 结果输出: 生成包含订单、车辆配置和路径的Excel文件

依赖库:
- pandas: 数据处理
- numpy: 数值计算
- openpyxl: Excel文件操作
- xgboost: 梯度提升树模型(需求预测)
- scikit-learn: 数据预处理和模型评估
- tensorflow: 神经网络模型(可选)
"""
import pandas as pd
import numpy as np
import openpyxl
from openpyxl import load_workbook
import warnings
from datetime import datetime, timedelta
import os
import logging
from typing import Dict, List, Tuple, Optional
import math
from pathlib import Path

try:
    import xgboost as xgb
    XGB_AVAILABLE = True
except ImportError:
    print("XGBoost is not installed; demand prediction will use the simple fallback.")
    XGB_AVAILABLE = False

try:
    from sklearn.preprocessing import LabelEncoder, StandardScaler
    from sklearn.model_selection import train_test_split
    from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score
    SKLEARN_AVAILABLE = True
except ImportError:
    print("scikit-learn is not installed; ML demand prediction and NN scaling will use fallbacks.")
    SKLEARN_AVAILABLE = False

# 尝试导入TensorFlow，用于神经网络模型
try:
    import tensorflow as tf
    from tensorflow import keras
    from tensorflow.keras import layers
    NN_AVAILABLE = True
except ImportError:
    print("TensorFlow未安装，将跳过神经网络功能")
    NN_AVAILABLE = False

# 配置日志
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

# 忽略警告信息
warnings.filterwarnings('ignore')

class BSSDemandPredictor:
    """自行车共享系统需求预测器
    
    该类负责基于历史骑行数据预测自行车站点的需求，使用XGBoost模型进行预测。
    主要功能包括数据加载、特征工程、模型训练和需求预测。
    """
    
    def __init__(self):
        # XGBoost回归模型实例
        self.xgb_model = None
        # 用于编码分类特征的标签编码器字典
        self.label_encoders = {}
        # 选择的目标自行车站点列表
        self.selected_stations = [
            "Pier 61 at Chelsea Piers",
            "West Drive & Prospect Park West", 
            "E 33 St & 1 Ave",
            "Cleveland Pl & Spring St",
            "N 7 St & Driggs Ave",
            "West End Ave & W 60 St",
            "N 6 St & Bedford Ave",
            "Hanson Pl & Ashland Pl",
            "Bergen St & Vanderbilt Ave",
            "Franklin St & Dupont St"
        ]
        
    def load_and_combine_data(self, csv_files: List[str]) -> pd.DataFrame:
        """加载并合并多个CSV文件的数据
        
        参数:
            csv_files: CSV文件路径列表，包含骑行数据
            
        返回:
            合并后的DataFrame，包含筛选和预处理后的数据
            
        功能:
            1. 逐块读取CSV文件，避免内存溢出
            2. 筛选出与目标站点相关的数据
            3. 将时间列转换为datetime类型并排序
            4. 提取时间特征(小时、星期几、日期)
        """
        logger.info("开始加载和合并CSV文件...")
        
        dataframes = []
        for i, csv_file in enumerate(csv_files, 1):
            logger.info(f"加载文件 {i}/{len(csv_files)}: {os.path.basename(csv_file)}")
            
            # 分块读取大文件，避免内存溢出
            chunk_size = 10000
            chunks = []
            
            for chunk in pd.read_csv(csv_file, chunksize=chunk_size):
                # 筛选包含目标站点的骑行记录
                chunk_filtered = chunk[
                    (chunk['start_station_name'].isin(self.selected_stations)) |
                    (chunk['end_station_name'].isin(self.selected_stations))
                ]
                if not chunk_filtered.empty:
                    chunks.append(chunk_filtered)
            
            if chunks:
                df = pd.concat(chunks, ignore_index=True)
                dataframes.append(df)
                logger.info(f"文件 {i} 加载完成，过滤后数据量: {len(df)}")
        
        if not dataframes:
            raise ValueError("没有找到相关站点的数据")
        
        # 合并所有数据
        combined_df = pd.concat(dataframes, ignore_index=True)
        logger.info(f"合并完成，总数据量: {len(combined_df)}")
        
        # 转换时间列格式
        combined_df['started_at'] = pd.to_datetime(combined_df['started_at'])
        combined_df['ended_at'] = pd.to_datetime(combined_df['ended_at'])
        
        # 按开始时间排序
        combined_df = combined_df.sort_values('started_at').reset_index(drop=True)
        
        # 提取时间特征
        combined_df['hour'] = combined_df['started_at'].dt.hour
        combined_df['day_of_week'] = combined_df['started_at'].dt.dayofweek
        combined_df['day'] = combined_df['started_at'].dt.date
        
        return combined_df
    
    def prepare_time_series_split(self, data: pd.DataFrame, train_days: int = 12, test_days: int = 3) -> Tuple[pd.DataFrame, pd.DataFrame]:
        logger.info(f"准备时间序列划分 - 训练集: {train_days}天, 测试集: {test_days}天")
        
        min_date = data['day'].min()
        max_date = data['day'].max()
        total_days = (max_date - min_date).days + 1
        
        logger.info(f"数据日期范围: {min_date} 到 {max_date}, 总共 {total_days} 天")
        
        if total_days < train_days + test_days:
            logger.warning(f"数据天数不足，调整为 {total_days-test_days} 天训练，{test_days} 天测试")
            train_days = total_days - test_days
        
        train_end_date = min_date + timedelta(days=train_days-1)
        test_start_date = train_end_date + timedelta(days=1)
        
        train_data = data[data['day'] <= train_end_date].copy()
        test_data = data[data['day'] >= test_start_date].copy()
        
        logger.info(f"训练集: {len(train_data)} 条记录 ({train_data['day'].min()} 到 {train_data['day'].max()})")
        logger.info(f"测试集: {len(test_data)} 条记录 ({test_data['day'].min()} 到 {test_data['day'].max()})")
        
        return train_data, test_data
    
    def create_hourly_demand_features(self, data: pd.DataFrame) -> pd.DataFrame:
        logger.info("创建每小时需求特征...")
        
        demand_features = []
        
        for station in self.selected_stations:
            departure_data = data[data['start_station_name'] == station].copy()
            if not departure_data.empty:
                departure_hourly = departure_data.groupby(['day', 'hour']).size().reset_index(name='departure_demand')
                departure_hourly['station'] = station
                
                arrival_data = data[data['end_station_name'] == station].copy()
                arrival_hourly = arrival_data.groupby(['day', 'hour']).size().reset_index(name='arrival_demand')
                
                hourly_demand = departure_hourly.merge(
                    arrival_hourly[['day', 'hour', 'arrival_demand']], 
                    on=['day', 'hour'], 
                    how='outer'
                ).fillna(0)
                
                hourly_demand['net_demand'] = hourly_demand['departure_demand'] - hourly_demand['arrival_demand']
                hourly_demand['station'] = station
                
                demand_features.append(hourly_demand)
        
        if demand_features:
            result = pd.concat(demand_features, ignore_index=True)
            result['day_of_week'] = pd.to_datetime(result['day']).dt.dayofweek
            result['is_weekend'] = result['day_of_week'].isin([5, 6]).astype(int)
            
            logger.info(f"需求特征创建完成，数据量: {len(result)}")
            return result
        else:
            logger.warning("没有创建需求特征")
            return pd.DataFrame()
    
    def train_model(self, train_data: pd.DataFrame) -> bool:
        logger.info("开始训练需求预测模型...")
        
        if not XGB_AVAILABLE or not SKLEARN_AVAILABLE:
            logger.warning("XGBoost or scikit-learn is unavailable; skipping model training.")
            return False

        train_features = self.create_hourly_demand_features(train_data)
        
        if train_features.empty or len(train_features) < 10:
            logger.warning("训练数据不足，跳过模型训练")
            return False
        
        feature_columns = ['hour', 'day_of_week', 'is_weekend', 'departure_demand', 'arrival_demand']
        
        if 'station' not in self.label_encoders:
            self.label_encoders['station'] = LabelEncoder()
            train_features['station_encoded'] = self.label_encoders['station'].fit_transform(train_features['station'])
        else:
            train_features['station_encoded'] = self.label_encoders['station'].transform(train_features['station'])
        
        feature_columns.append('station_encoded')
        
        X = train_features[feature_columns].fillna(0)
        y = train_features['net_demand'].fillna(0)
        
        try:
            self.xgb_model = xgb.XGBRegressor(
                n_estimators=100,
                max_depth=6,
                learning_rate=0.1,
                random_state=42,
                n_jobs=-1
            )
            
            self.xgb_model.fit(X, y)
            
            y_pred = self.xgb_model.predict(X)
            mae = mean_absolute_error(y, y_pred)
            rmse = np.sqrt(mean_squared_error(y, y_pred))
            r2 = r2_score(y, y_pred)
            
            logger.info(f"模型训练完成 - MAE: {mae:.2f}, RMSE: {rmse:.2f}, R²: {r2:.3f}")
            return True
            
        except Exception as e:
            logger.error(f"模型训练失败: {e}")
            return False
    
    def predict_demand(self, test_data: pd.DataFrame) -> pd.DataFrame:
        if self.xgb_model is None or not SKLEARN_AVAILABLE:
            logger.info("使用简单预测方法")
            return self.generate_simple_predictions(test_data)
        
        logger.info("使用训练好的模型进行需求预测...")
        
        test_features = self.create_hourly_demand_features(test_data)
        
        if test_features.empty:
            logger.warning("测试数据为空，使用简单预测")
            return self.generate_simple_predictions(test_data)
        
        feature_columns = ['hour', 'day_of_week', 'is_weekend', 'departure_demand', 'arrival_demand']
        test_features['station_encoded'] = self.label_encoders['station'].transform(test_features['station'])
        feature_columns.append('station_encoded')
        
        X_test = test_features[feature_columns].fillna(0)
        predictions = self.xgb_model.predict(X_test)
        test_features['predicted_net_demand'] = predictions
        
        logger.info(f"预测完成，预测数据量: {len(test_features)}")
        return test_features
    
    def generate_simple_predictions(self, test_data: pd.DataFrame) -> pd.DataFrame:
        logger.info("生成简单预测结果...")
        
        predictions = []
        dates = test_data['day'].unique()
        
        for station in self.selected_stations:
            for date in dates:
                for hour in range(24):
                    if hour in [7, 8, 9, 17, 18, 19]:
                        predicted_demand = np.random.randint(-20, 30)
                    else:
                        predicted_demand = np.random.randint(-10, 15)
                    
                    predictions.append({
                        'station': station,
                        'day': date,
                        'hour': hour,
                        'predicted_net_demand': predicted_demand,
                        'day_of_week': pd.to_datetime(date).dayofweek
                    })
        
        result = pd.DataFrame(predictions)
        logger.info(f"简单预测完成，预测数据量: {len(result)}")
        return result

class BSSFleetOptimizer:
    
    def __init__(self, use_neural_network: bool = False, predefined_vehicles: List[Dict] = None):
        self.use_neural_network = use_neural_network and NN_AVAILABLE and SKLEARN_AVAILABLE
        self.scaler = StandardScaler() if SKLEARN_AVAILABLE else None
        self.nn_model = None
        
        if predefined_vehicles:
            self.vehicle_types = predefined_vehicles
            logger.info(f"使用预定义载具: {len(self.vehicle_types)}种")
        else:
            self.vehicle_types = [
                {'name': 'Barge1', 'u': 30, 'speed': 45, 'c1': 0.8, 'c1_prime': 0.3, 'c4': 0.2, 'type': 1},
                {'name': 'Barge2', 'u': 30, 'speed': 45, 'c1': 0.8, 'c1_prime': 0.3, 'c4': 0.2, 'type': 1},
                {'name': 'Barge3', 'u': 30, 'speed': 45, 'c1': 0.8, 'c1_prime': 0.3, 'c4': 0.2, 'type': 1},
                {'name': 'Barge4', 'u': 30, 'speed': 45, 'c1': 0.8, 'c1_prime': 0.3, 'c4': 0.2, 'type': 1},
                {'name': 'Barge5', 'u': 30, 'speed': 45, 'c1': 0.8, 'c1_prime': 0.3, 'c4': 0.2, 'type': 1},
                {'name': 'Train1', 'u': 60, 'speed': 35, 'c1': 0.85, 'c1_prime': 0.4, 'c4': 0.15, 'type': 2},
                {'name': 'Train2', 'u': 60, 'speed': 35, 'c1': 0.85, 'c1_prime': 0.4, 'c4': 0.15, 'type': 2},
                {'name': 'Train3', 'u': 60, 'speed': 35, 'c1': 0.85, 'c1_prime': 0.4, 'c4': 0.15, 'type': 2},
                {'name': 'Train4', 'u': 60, 'speed': 35, 'c1': 0.85, 'c1_prime': 0.4, 'c4': 0.15, 'type': 2},
                {'name': 'Train5', 'u': 60, 'speed': 35, 'c1': 0.85, 'c1_prime': 0.4, 'c4': 0.15, 'type': 2},
                {'name': 'Truck1', 'u': 10, 'speed': 60, 'c1': 0.75, 'c1_prime': 0.25, 'c4': 0.1, 'type': 3},
                {'name': 'Truck2', 'u': 10, 'speed': 60, 'c1': 0.75, 'c1_prime': 0.25, 'c4': 0.1, 'type': 3},
                {'name': 'Truck3', 'u': 10, 'speed': 60, 'c1': 0.75, 'c1_prime': 0.25, 'c4': 0.1, 'type': 3},
                {'name': 'Truck4', 'u': 10, 'speed': 60, 'c1': 0.75, 'c1_prime': 0.25, 'c4': 0.1, 'type': 3},
                {'name': 'Truck5', 'u': 10, 'speed': 60, 'c1': 0.75, 'c1_prime': 0.25, 'c4': 0.1, 'type': 3}
            ]
    
    def create_neural_network_model(self, input_dim: int) -> None:
        if not self.use_neural_network:
            return
            
        logger.info("创建神经网络配车模型...")
        
        model = keras.Sequential([
            layers.Dense(64, activation='relu', input_shape=(input_dim,)),
            layers.Dropout(0.2),
            layers.Dense(32, activation='relu'),
            layers.Dropout(0.2),
            layers.Dense(16, activation='relu'),
            layers.Dense(len(self.vehicle_types), activation='softmax')
        ])
        
        model.compile(
            optimizer='adam',
            loss='categorical_crossentropy',
            metrics=['accuracy']
        )
        
        self.nn_model = model
        logger.info("神经网络模型创建完成")
    
    def prepare_training_data(self, orders: List[Dict]) -> Tuple[np.ndarray, np.ndarray]:
        if not orders:
            return np.array([]), np.array([])
        
        features = []
        labels = []
        
        for order in orders:
            feature = [
                order['qr'],
                order['ap'],
                len(order.get('p', '')),
                len(order.get('d', '')),
                abs(hash(order.get('p', '')) % 100),
                abs(hash(order.get('d', '')) % 100),
            ]
            
            best_vehicle_idx = self.select_vehicle_traditional(order)
            label = [0] * len(self.vehicle_types)
            label[best_vehicle_idx] = 1
            
            features.append(feature)
            labels.append(label)
        
        return np.array(features), np.array(labels)
    
    def select_vehicle_traditional(self, order: Dict) -> int:
        demand = order['qr']
        
        suitable_vehicles = []
        for i, vehicle in enumerate(self.vehicle_types):
            if vehicle['u'] >= demand:
                utilization = demand / vehicle['u']
                cost_factor = vehicle['c1'] + vehicle['c1_prime'] + vehicle['c4']
                efficiency = utilization / cost_factor
                suitable_vehicles.append((i, efficiency))
        
        if suitable_vehicles:
            return max(suitable_vehicles, key=lambda x: x[1])[0]
        else:
            return max(enumerate(self.vehicle_types), key=lambda x: x[1]['u'])[0]
    
    def train_neural_network(self, orders: List[Dict]) -> bool:
        if not self.use_neural_network or not orders:
            return False
        
        logger.info("开始训练神经网络配车模型...")
        
        X, y = self.prepare_training_data(orders)
        
        if len(X) == 0:
            logger.warning("没有训练数据")
            return False
        
        X_scaled = self.scaler.fit_transform(X)
        
        self.create_neural_network_model(X.shape[1])
        
        try:
            self.nn_model.fit(
                X_scaled, y,
                epochs=50,
                batch_size=32,
                validation_split=0.2,
                verbose=0
            )
            logger.info("神经网络训练完成")
            return True
        except Exception as e:
            logger.error(f"神经网络训练失败: {e}")
            return False
    
    def select_vehicles_neural_network(self, orders: List[Dict]) -> pd.DataFrame:
        if not self.use_neural_network or self.nn_model is None:
            return self.select_vehicles_traditional_method(orders)
        
        logger.info("使用神经网络进行车辆配置...")
        
        features = []
        for order in orders:
            feature = [
                order['qr'],
                order['ap'],
                len(order.get('p', '')),
                len(order.get('d', '')),
                abs(hash(order.get('p', '')) % 100),
                abs(hash(order.get('d', '')) % 100),
            ]
            features.append(feature)
        
        if not features:
            return self.select_vehicles_traditional_method(orders)
        
        X = np.array(features)
        X_scaled = self.scaler.transform(X)
        
        predictions = self.nn_model.predict(X_scaled, verbose=0)
        
        total_demand = sum(order['qr'] for order in orders)
        
        vehicle_demand_scores = np.sum(predictions, axis=0)
        
        selected_vehicles = []
        vehicle_id = 1
        remaining_demand = total_demand
        
        vehicle_priorities = sorted(enumerate(vehicle_demand_scores), key=lambda x: x[1], reverse=True)
        
        for vehicle_idx, demand_score in vehicle_priorities:
            if remaining_demand <= 0:
                break
            
            vehicle_type = self.vehicle_types[vehicle_idx]
            
            if demand_score > 0.1:
                max_vehicles_needed = max(1, int(np.ceil(remaining_demand / vehicle_type['u'])))
                score_based_count = max(1, int(np.ceil(demand_score)))
                vehicle_count = min(max_vehicles_needed, score_based_count, 5)
                
                for _ in range(vehicle_count):
                    if remaining_demand <= 0:
                        break
                    
                    vehicle = vehicle_type.copy()
                    vehicle['K'] = f"{vehicle['name']}{vehicle_id}"
                    selected_vehicles.append(vehicle)
                    remaining_demand -= vehicle['u']
                    vehicle_id += 1
        
        if remaining_demand > 0 and len(selected_vehicles) < 10:
            best_vehicle = max(self.vehicle_types, key=lambda x: x['u'] / (x['c1'] + x['c1_prime'] + x['c4']))
            while remaining_demand > 0 and len(selected_vehicles) < 10:
                vehicle = best_vehicle.copy()
                vehicle['K'] = f"{vehicle['name']}{vehicle_id}"
                selected_vehicles.append(vehicle)
                remaining_demand -= vehicle['u']
                vehicle_id += 1
        
        while len(selected_vehicles) < 3:
            vehicle = self.vehicle_types[0].copy()
            vehicle['K'] = f"{vehicle['name']}{vehicle_id}"
            selected_vehicles.append(vehicle)
            vehicle_id += 1
        
        result_df = pd.DataFrame(selected_vehicles)
        logger.info(f"神经网络配车完成，选择了 {len(result_df)} 辆车")
        logger.info(f"总需求: {total_demand}, 总容量: {sum(v['u'] for v in selected_vehicles)}")
        
        return result_df
    
    def select_vehicles_traditional_method(self, orders: List[Dict]) -> pd.DataFrame:
        logger.info("使用传统算法进行车辆配置...")
        
        if not orders:
            selected_vehicles = []
            for i, vehicle in enumerate(self.vehicle_types[:5]):
                vehicle_copy = vehicle.copy()
                vehicle_copy['K'] = f"{vehicle['name']}"
                selected_vehicles.append(vehicle_copy)
            
            return pd.DataFrame(selected_vehicles)
        
        total_demand = sum(order['qr'] for order in orders)
        max_single_demand = max([order['qr'] for order in orders])
        
        logger.info(f"订单分析 - 总需求: {total_demand}, 最大单次需求: {max_single_demand}")
        
        selected_vehicles = []
        vehicle_id = 1
        remaining_demand = total_demand
        
        while remaining_demand > 0 and len(selected_vehicles) < 20:
            best_vehicle = None
            best_efficiency = 0
            
            for vehicle in self.vehicle_types:
                if vehicle['u'] <= remaining_demand + 50:
                    cost_factor = vehicle['c1'] + vehicle['c1_prime'] + vehicle['c4']
                    efficiency = vehicle['u'] / cost_factor
                    if efficiency > best_efficiency:
                        best_efficiency = efficiency
                        best_vehicle = vehicle
            
            if best_vehicle is None:
                best_vehicle = min(self.vehicle_types, key=lambda x: x['u'])
            
            vehicle_copy = best_vehicle.copy()
            vehicle_copy['K'] = f"{best_vehicle['name']}{vehicle_id}"
            selected_vehicles.append(vehicle_copy)
            
            remaining_demand -= best_vehicle['u']
            vehicle_id += 1
        
        while len(selected_vehicles) < 5:
            vehicle = np.random.choice(self.vehicle_types)
            vehicle_copy = vehicle.copy()
            vehicle_copy['K'] = f"{vehicle['name']}{vehicle_id}"
            selected_vehicles.append(vehicle_copy)
            vehicle_id += 1
        
        result_df = pd.DataFrame(selected_vehicles)
        logger.info(f"传统配车完成，选择了 {len(result_df)} 辆车")
        
        return result_df

class BSSCompleteSystem:
    
    def __init__(self, config: Dict):
        self.config = config
        self.project_root = Path(config.get('project_root', Path(__file__).resolve().parents[2]))
        
        self.target_order_count = config.get('target_order_count', 20)
        self.random_seed = config.get('random_seed', 42)
        np.random.seed(self.random_seed)
        
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        self.result_dir = Path(config.get('result_dir', self.project_root / "outputs" / "generated_bss_inputs"))
        self.result_dir.mkdir(parents=True, exist_ok=True)
        self.traditional_excel_path = f"{self.result_dir}/Intermodal_EGS_data_all__traditional_{timestamp}.xlsx"
        self.neural_excel_path = f"{self.result_dir}/Intermodal_EGS_data_all__neural_{timestamp}.xlsx"
        
        self.predefined_vehicles_path = Path(config.get(
            'predefined_vehicles_path',
            self.project_root / "data" / "templates" / "Intermodal_EGS_data_all.xlsx"
        ))
        
        self.predefined_vehicles = self.load_predefined_vehicles()
        
        self.predictor = BSSDemandPredictor()
        self.optimizer_traditional = BSSFleetOptimizer(use_neural_network=False, predefined_vehicles=self.predefined_vehicles)
        self.optimizer_neural = BSSFleetOptimizer(use_neural_network=True, predefined_vehicles=self.predefined_vehicles)
        
        configured_csv_files = config.get('csv_files')
        if configured_csv_files:
            self.csv_files = [str(Path(p)) for p in configured_csv_files]
        else:
            sample_dir = self.project_root / "data" / "raw" / "citibike_sample"
            raw_dir = self.project_root / "data" / "raw" / "citibike"
            self.csv_files = [str(p) for p in sorted(sample_dir.glob("*.csv")) + sorted(raw_dir.glob("*.csv"))]
        
        self.base_excel_path = Path(config.get(
            'base_excel_path',
            self.project_root / "data" / "templates" / "Intermodal_EGS_data_all.xlsx"
        ))
    
    def load_predefined_vehicles(self) -> List[Dict]:
        logger.info("加载预定义载具列表...")
        
        try:
            wb = openpyxl.load_workbook(self.predefined_vehicles_path)
            ws = wb['K']
            
            vehicles = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0]:
                    vehicle = {
                        'name': row[0],
                        'u': row[1],
                        'speed': row[2],
                        'c1': row[3],
                        'c1_prime': row[4],
                        'c4': row[5],
                        'type': row[6]
                    }
                    vehicles.append(vehicle)
            
            logger.info(f"加载了 {len(vehicles)} 种载具")
            logger.info(f"载具类型分布: Type 1: {len([v for v in vehicles if v['type'] == 1])}个, "
                       f"Type 2: {len([v for v in vehicles if v['type'] == 2])}个, "
                       f"Type 3: {len([v for v in vehicles if v['type'] == 3])}个")
            
            return vehicles
            
        except Exception as e:
            logger.error(f"加载预定义载具失败: {e}")
            return [
                {'name': 'Barge1', 'u': 30, 'speed': 45, 'c1': 0.08, 'c1_prime': 0.15, 'c4': 0.2, 'type': 1},
                {'name': 'Train1', 'u': 60, 'speed': 35, 'c1': 0.10, 'c1_prime': 0.2, 'c4': 0.15, 'type': 2},
                {'name': 'Truck1', 'u': 10, 'speed': 60, 'c1': 0.075, 'c1_prime': 0.125, 'c4': 0.1, 'type': 3}
            ]
    
    def clean_location_name(self, name):
        if isinstance(name, str):
            cleaned = name.replace('&', '')
            cleaned = ' '.join(cleaned.split())
            return cleaned
        return name
    
    def convert_predictions_to_orders(self, predictions: pd.DataFrame) -> List[Dict]:
        logger.info("将预测结果转换为订单...")
        
        orders = []
        order_id = 1
        
        for _, row in predictions.iterrows():
            demand = int(row['predicted_net_demand'])
            
            if abs(demand) < 5:
                continue
            
            station = self.clean_location_name(row['station'])
            hour = int(row['hour'])
            
            if demand > 0:
                target_stations = [self.clean_location_name(s) for s in self.predictor.selected_stations if s != row['station']]
                if target_stations:
                    target = np.random.choice(target_stations)
                    
                    order = {
                        'order_id': order_id,
                        'p': station,
                        'd': target,
                        'ap': hour,
                        'bp': hour + 1,
                        'ad': hour,
                        'bd': hour + 1,
                        'qr': min(abs(demand), 40)
                    }
                    orders.append(order)
                    order_id += 1
        
        logger.info(f"生成 {len(orders)} 个订单")
        return orders
    
    def distribute_orders_to_sheets(self, orders: List[Dict]) -> Dict[str, List[Dict]]:
        logger.info("分配订单到不同的sheets...")
        
        distributed_orders = {}
        target_counts = [10, 20, 50, 100]
        
        distributed_orders['R'] = orders
        logger.info(f"R: 分配了 {len(orders)} 个订单")
        
        for count in target_counts:
            if count <= len(orders):
                selected_orders = orders[:count]
                distributed_orders[f'R_{count}'] = selected_orders
                logger.info(f"R_{count}: 分配了 {len(selected_orders)} 个订单")
        
        if self.target_order_count not in target_counts and self.target_order_count <= len(orders):
            selected_orders = orders[:self.target_order_count]
            distributed_orders[f'R_{self.target_order_count}'] = selected_orders
            logger.info(f"R_{self.target_order_count}: 分配了 {len(selected_orders)} 个订单（用户指定）")
        
        return distributed_orders
    
    def get_target_orders(self, distributed_orders: Dict[str, List[Dict]]) -> List[Dict]:
        target_sheet = f'R_{self.target_order_count}'
        
        if target_sheet in distributed_orders:
            target_orders = distributed_orders[target_sheet]
            logger.info(f"使用 {target_sheet} 的 {len(target_orders)} 个订单进行配车算法")
            return target_orders
        else:
            all_orders = distributed_orders.get('R', [])
            actual_count = min(self.target_order_count, len(all_orders))
            target_orders = all_orders[:actual_count]
            logger.info(f"订单数量不足，使用前 {len(target_orders)} 个订单进行配车算法")
            return target_orders
    
    def create_o_sheet_data(self, orders: List[Dict], vehicle_config: pd.DataFrame) -> pd.DataFrame:
        logger.info("创建o表数据...")
        
        o_data = []
        
        for _, vehicle in vehicle_config.iterrows():
            if orders:
                assigned_order = np.random.choice(orders)
                o_data.append({
                    'K': vehicle['K'],
                    'o': assigned_order['p'],
                    'o2': assigned_order['d']
                })
            else:
                o_data.append({
                    'K': vehicle['K'],
                    'o': 'N/A',
                    'o2': 'N/A'
                })
        
        result_df = pd.DataFrame(o_data)
        logger.info(f"o表创建完成，数据量: {len(result_df)}")
        return result_df
    
    def copy_original_sheets(self, workbook, original_file_path):
        try:
            original_wb = openpyxl.load_workbook(original_file_path)
            
            if 'N' in original_wb.sheetnames:
                original_n = original_wb['N']
                new_n = workbook.create_sheet('N')
                for row in original_n.iter_rows():
                    for cell in row:
                        if cell.value is not None:
                            value = self.clean_location_name(cell.value) if isinstance(cell.value, str) else cell.value
                            new_n.cell(row=cell.row, column=cell.column, value=value)
                logger.info("已复制N sheet")
            
            if 'T' in original_wb.sheetnames:
                original_t = original_wb['T']
                new_t = workbook.create_sheet('T')
                for row in original_t.iter_rows():
                    for cell in row:
                        if cell.value is not None:
                            value = self.clean_location_name(cell.value) if isinstance(cell.value, str) else cell.value
                            new_t.cell(row=cell.row, column=cell.column, value=value)
                logger.info("已复制T sheet")
                
        except Exception as e:
            logger.warning(f"复制原始sheets失败: {e}")
    
    def create_dual_excel_files(self, distributed_orders: Dict[str, List[Dict]], vehicle_configs: Dict[str, pd.DataFrame], o_sheets: Dict[str, pd.DataFrame]):
        logger.info("开始创建两个分离的Excel文件...")
        
        try:
            traditional_configs = {k: v for k, v in vehicle_configs.items() if 'traditional' in k}
            neural_configs = {k: v for k, v in vehicle_configs.items() if 'neural' in k}
            
            traditional_o_sheets = {k: v for k, v in o_sheets.items() if 'traditional' in k}
            neural_o_sheets = {k: v for k, v in o_sheets.items() if 'neural' in k}

            created_files = []
            if traditional_configs:
                self.create_single_excel_file(
                    self.traditional_excel_path,
                    self.base_excel_path,
                    distributed_orders,
                    traditional_configs,
                    traditional_o_sheets,
                    "traditional algorithm"
                )
                created_files.append(("traditional", self.traditional_excel_path))

            if neural_configs:
                self.create_single_excel_file(
                    self.neural_excel_path,
                    self.base_excel_path,
                    distributed_orders,
                    neural_configs,
                    neural_o_sheets,
                    "neural algorithm"
                )
                created_files.append(("neural", self.neural_excel_path))

            logger.info("Excel files created:")
            for algorithm_name, file_path in created_files:
                logger.info(f"  {algorithm_name}: {file_path}")
            return
            
            self.create_single_excel_file(
                self.traditional_excel_path, 
                self.base_excel_path,
                distributed_orders, 
                traditional_configs, 
                traditional_o_sheets, 
                "传统算法"
            )
            
            self.create_single_excel_file(
                self.neural_excel_path, 
                self.base_excel_path,
                distributed_orders, 
                neural_configs, 
                neural_o_sheets, 
                "神经网络算法"
            )
            
            logger.info(f"两个Excel文件创建完成:")
            logger.info(f"  传统算法: {self.traditional_excel_path}")
            logger.info(f"  神经网络: {self.neural_excel_path}")
            
        except Exception as e:
            logger.error(f"创建Excel文件失败: {e}")
            raise
    
    def create_single_excel_file(self, file_path: str, original_file_path: str, distributed_orders: Dict[str, List[Dict]], 
                                 vehicle_configs: Dict[str, pd.DataFrame], o_sheets: Dict[str, pd.DataFrame], algorithm_name: str):
        workbook = openpyxl.Workbook()
        
        default_sheet = workbook.active
        workbook.remove(default_sheet)
        
        self.copy_original_sheets(workbook, original_file_path)
        
        for sheet_name, orders in distributed_orders.items():
            sheet = workbook.create_sheet(sheet_name)
            
            headers = ['p', 'd', 'ap', 'bp', 'ad', 'bd', 'qr']
            for col, header in enumerate(headers, 1):
                sheet.cell(row=1, column=col, value=header)
            
            for row, order in enumerate(orders, 2):
                sheet.cell(row=row, column=1, value=order['p'])
                sheet.cell(row=row, column=2, value=order['d'])
                sheet.cell(row=row, column=3, value=order['ap'])
                sheet.cell(row=row, column=4, value=order['bp'])
                sheet.cell(row=row, column=5, value=order['ad'])
                sheet.cell(row=row, column=6, value=order['bd'])
                sheet.cell(row=row, column=7, value=order['qr'])
            
            logger.info(f"({algorithm_name}) 订单数据写入 {sheet_name} sheet，共 {len(orders)} 条记录")
        
        for config_name, vehicle_config in vehicle_configs.items():
            sheet_name = 'K'
            sheet = workbook.create_sheet(sheet_name)
            
            headers = ['K', 'u', 'speed', 'c1', 'c1\'', 'c4', 'type']
            for col, header in enumerate(headers, 1):
                sheet.cell(row=1, column=col, value=header)
            
            for row, (_, vehicle) in enumerate(vehicle_config.iterrows(), 2):
                sheet.cell(row=row, column=1, value=vehicle['K'])
                sheet.cell(row=row, column=2, value=vehicle['u'])
                sheet.cell(row=row, column=3, value=vehicle['speed'])
                sheet.cell(row=row, column=4, value=vehicle['c1'])
                sheet.cell(row=row, column=5, value=vehicle['c1_prime'])
                sheet.cell(row=row, column=6, value=vehicle['c4'])
                sheet.cell(row=row, column=7, value=vehicle['type'])
            
            logger.info(f"({algorithm_name}) 车辆配置写入 K sheet，共 {len(vehicle_config)} 辆车")
            break
        
        for o_name, o_data in o_sheets.items():
            sheet_name = 'o'
            sheet = workbook.create_sheet(sheet_name)
            
            headers = ['K', 'o', 'o2']
            for col, header in enumerate(headers, 1):
                sheet.cell(row=1, column=col, value=header)
            
            for row, (_, path) in enumerate(o_data.iterrows(), 2):
                sheet.cell(row=row, column=1, value=path['K'])
                sheet.cell(row=row, column=2, value=path['o'])
                sheet.cell(row=row, column=3, value=path['o2'])
            
            logger.info(f"({algorithm_name}) 路径信息写入 o sheet，共 {len(o_data)} 条记录")
            break
        
        workbook.save(file_path)
        logger.info(f"({algorithm_name}) Excel文件保存至: {file_path}")

    def update_excel_file(self, distributed_orders: Dict[str, List[Dict]], vehicle_configs: Dict[str, pd.DataFrame], o_sheets: Dict[str, pd.DataFrame]):
        self.create_dual_excel_files(distributed_orders, vehicle_configs, o_sheets)
    
    def run_system(self):
        logger.info("=" * 50)
        logger.info("启动BSS完整系统")
        logger.info(f"配置: {self.config}")
        logger.info("=" * 50)
        
        try:
            combined_data = self.predictor.load_and_combine_data(self.csv_files)
            train_data, test_data = self.predictor.prepare_time_series_split(combined_data)
            
            if self.config.get('enable_demand_prediction', True):
                model_trained = self.predictor.train_model(train_data)
                predictions = self.predictor.predict_demand(test_data)
            else:
                logger.info("需求预测模块已禁用，使用简单预测")
                predictions = self.predictor.generate_simple_predictions(test_data)
            
            orders = self.convert_predictions_to_orders(predictions)
            
            if not orders:
                logger.warning("没有生成订单，创建示例订单")
                orders = [
                    {'order_id': 1, 'p': self.predictor.selected_stations[0], 'd': self.predictor.selected_stations[1], 
                     'ap': 12, 'bp': 13, 'ad': 12, 'bd': 13, 'qr': 20},
                    {'order_id': 2, 'p': self.predictor.selected_stations[2], 'd': self.predictor.selected_stations[3], 
                     'ap': 15, 'bp': 16, 'ad': 15, 'bd': 16, 'qr': 30}
                ]
            
            distributed_orders = self.distribute_orders_to_sheets(orders)
            
            target_orders = self.get_target_orders(distributed_orders)
            
            vehicle_configs = {}
            o_sheets = {}
            
            if self.config.get('enable_traditional_fleet', True):
                traditional_config = self.optimizer_traditional.select_vehicles_traditional_method(target_orders)
                vehicle_configs[f'K_traditional_R_{self.target_order_count}'] = traditional_config
                o_sheets[f'o_traditional_R_{self.target_order_count}'] = self.create_o_sheet_data(target_orders, traditional_config)
            
            if self.config.get('enable_neural_fleet', True):
                nn_trained = self.optimizer_neural.train_neural_network(target_orders)
                neural_config = self.optimizer_neural.select_vehicles_neural_network(target_orders)
                vehicle_configs[f'K_neural_R_{self.target_order_count}'] = neural_config
                o_sheets[f'o_neural_R_{self.target_order_count}'] = self.create_o_sheet_data(target_orders, neural_config)
            
            if not vehicle_configs:
                default_config = self.optimizer_traditional.select_vehicles_traditional_method([])
                vehicle_configs['K'] = default_config
                o_sheets['o'] = self.create_o_sheet_data([], default_config)
            
            self.create_dual_excel_files(distributed_orders, vehicle_configs, o_sheets)
            
            logger.info("=" * 50)
            logger.info("系统运行完成")
            logger.info(f"生成订单总数: {len(orders)}")
            logger.info(f"用户指定订单数: {self.target_order_count}")
            logger.info(f"分配sheets: {list(distributed_orders.keys())}")
            logger.info(f"车队配置: {list(vehicle_configs.keys())}")
            logger.info(f"配置传统算法: {'是' if self.config.get('enable_traditional_fleet') else '否'}")
            logger.info(f"配置神经网络: {'是' if self.config.get('enable_neural_fleet') else '否'}")
            
            target_orders = self.get_target_orders(distributed_orders)
            total_demand = sum(order['qr'] for order in target_orders)
            logger.info(f"目标订单总需求: {total_demand}")
            
            for config_name, vehicle_config in vehicle_configs.items():
                total_capacity = sum(vehicle_config['u']) if len(vehicle_config) > 0 else 0
                vehicle_count = len(vehicle_config)
                utilization = (total_demand / total_capacity * 100) if total_capacity > 0 else 0
                logger.info(f"{config_name}: {vehicle_count}辆车, 总容量{total_capacity}, 利用率{utilization:.1f}%")
            
            logger.info("=" * 50)
            
        except Exception as e:
            logger.error(f"系统运行失败: {e}")
            raise

def main():
    import sys
    
    target_order_count = 20
    if len(sys.argv) > 1:
        try:
            target_order_count = int(sys.argv[1])
            logger.info(f"用户指定订单数量: {target_order_count}")
        except ValueError:
            logger.warning(f"无效的订单数量参数: {sys.argv[1]}，使用默认值: {target_order_count}")
    
    config = {
        'enable_demand_prediction': True,
        'enable_traditional_fleet': True,
        'enable_neural_fleet': True,
        'target_order_count': target_order_count,
        'name': f'完整BSS系统(R_{target_order_count})'
    }
    
    logger.info(f"启动BSS完整系统，目标订单数: {target_order_count}...")
    system = BSSCompleteSystem(config)
    system.run_system()

def cli_main():
    import argparse

    default_root = Path(__file__).resolve().parents[2]
    parser = argparse.ArgumentParser(
        description="Run the integrated BSS demand prediction and fleet configuration pipeline."
    )
    parser.add_argument(
        "target_order_count",
        nargs="?",
        type=int,
        default=20,
        help="Number of orders used by the fleet configuration step, for example 10, 20, 50, or 100.",
    )
    parser.add_argument(
        "--project-root",
        default=str(default_root),
        help="Repository root. Defaults to the current repository inferred from this script.",
    )
    parser.add_argument(
        "--csv-dir",
        action="append",
        help="Directory containing CitiBike CSV files. Can be passed more than once.",
    )
    parser.add_argument(
        "--csv-file",
        action="append",
        help="Single CitiBike CSV file. Can be passed more than once.",
    )
    parser.add_argument("--result-dir", help="Directory for generated Excel files.")
    parser.add_argument("--template", help="Base Excel template containing N/T sheets.")
    parser.add_argument("--vehicles-template", help="Excel template containing the K vehicle sheet.")
    parser.add_argument(
        "--disable-demand-prediction",
        action="store_true",
        help="Skip XGBoost training and use simple random demand generation.",
    )
    parser.add_argument("--no-traditional", action="store_true", help="Skip traditional vehicle selection output.")
    parser.add_argument("--no-neural", action="store_true", help="Skip neural-network vehicle selection output.")
    parser.add_argument("--seed", type=int, default=42, help="Random seed for reproducible generated data.")
    args = parser.parse_args()

    project_root = Path(args.project_root).resolve()
    csv_files = []
    for csv_dir in args.csv_dir or []:
        csv_files.extend(sorted(Path(csv_dir).glob("*.csv")))
    csv_files.extend(Path(p) for p in args.csv_file or [])

    config = {
        'enable_demand_prediction': not args.disable_demand_prediction,
        'enable_traditional_fleet': not args.no_traditional,
        'enable_neural_fleet': not args.no_neural,
        'target_order_count': args.target_order_count,
        'random_seed': args.seed,
        'project_root': project_root,
        'name': f'Complete BSS system (R_{args.target_order_count})',
    }
    if csv_files:
        config['csv_files'] = csv_files
    if args.result_dir:
        config['result_dir'] = Path(args.result_dir)
    if args.template:
        config['base_excel_path'] = Path(args.template)
    if args.vehicles_template:
        config['predefined_vehicles_path'] = Path(args.vehicles_template)

    logger.info(f"Starting BSS complete system; target orders: {args.target_order_count}")
    system = BSSCompleteSystem(config)
    system.run_system()


if __name__ == "__main__":
    cli_main()
